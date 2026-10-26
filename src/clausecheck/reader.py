"""Turn an uploaded contract (PDF, DOCX or plain text) into plain text."""

from __future__ import annotations

import io
from pathlib import Path


def read_bytes(data: bytes, filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if suffix == ".docx":
        import docx

        doc = docx.Document(io.BytesIO(data))
        lines = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(lines)
    if suffix in (".txt", ".md", ""):
        for enc in ("utf-8", "cp1252", "latin-1"):
            try:
                return data.decode(enc)
            except UnicodeDecodeError:
                continue
    raise ValueError(f"Unsupported file type: {suffix or 'unknown'}. Use PDF, DOCX or TXT.")


def read_file(path: str | Path) -> str:
    path = Path(path)
    return read_bytes(path.read_bytes(), path.name)
