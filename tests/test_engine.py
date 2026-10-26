"""Unit tests for segmentation, checks, negation handling and rule validation."""

import copy

import pytest
import yaml

from clausecheck import Rules, analyse, read_bytes, to_csv, to_markdown
from clausecheck.engine import RULES_DIR, segment

RULES = Rules.builtin()


def ids(text):
    return {f.id for f in analyse(text, RULES).findings}


def test_segment_numbered_headings():
    segs = segment("1. Term\nRuns for 12 months.\n2. Termination\nEither party may terminate.")
    assert [s.heading for s in segs] == ["1. Term", "2. Termination"]


def test_numbered_sentence_is_not_a_heading():
    segs = segment("1. Fees\n1.1 The Customer shall pay all invoices within thirty days of receipt.\n")
    assert len(segs) == 1


@pytest.mark.parametrize("sentence", [
    "The Supplier may use Customer Data to train and improve its models.",
    "Provider may analyse your data to develop new AI services.",
    "Leverandøren må anvende Kundens data til at træne og forbedre modeller.",
])
def test_training_on_data_is_flagged(sentence):
    assert "ai_training_on_data" in ids(sentence)


@pytest.mark.parametrize("sentence", [
    "The Supplier shall not use Customer Data to train its models.",
    "The Supplier will never use your data to improve its models.",
    "The Supplier may use Customer Data to improve its models only with the Customer's prior written consent.",
    "Leverandøren må ikke anvende Kundens data til at træne modeller.",
])
def test_negated_or_consent_based_training_is_not_flagged(sentence):
    assert "ai_training_on_data" not in ids(sentence)


def test_personal_data_without_dpa_is_flagged_and_dpa_clears_it():
    assert "no_dpa" in ids("The vendor processes personal data of employees.")
    assert "no_dpa" not in ids("The vendor processes personal data under a data processing agreement (Article 28 GDPR).")


def test_eu_law_is_fine_us_law_is_flagged():
    assert "non_eu_law" not in ids("9. Governing Law\nThis Agreement is governed by the laws of Denmark.")
    assert "non_eu_law" in ids("9. Governing Law\nThis Agreement is governed by the laws of the State of California.")


def test_reader_rejects_unknown_types():
    with pytest.raises(ValueError):
        read_bytes(b"x", "contract.xlsx")


def test_reader_handles_docx(tmp_path):
    import docx

    d = docx.Document()
    d.add_paragraph("1. Governing Law")
    d.add_paragraph("This Agreement is governed by the laws of Denmark.")
    path = tmp_path / "c.docx"
    d.save(path)
    text = read_bytes(path.read_bytes(), "c.docx")
    assert analyse(text, RULES).values["governing_law"] == "Denmark"


@pytest.mark.parametrize("mutate, msg", [
    (lambda c, k: k["checks"].append(dict(k["checks"][0])), "Duplicate"),
    (lambda c, k: k["checks"][0].update(severity="critical"), "severity"),
    (lambda c, k: k["checks"][0].update(why=""), "missing"),
    (lambda c, k: next(x for x in k["checks"] if x["type"] == "missing").update(clause="nope"), "unknown clause"),
])
def test_rule_validation(mutate, msg):
    c = yaml.safe_load((RULES_DIR / "clauses.yaml").read_text(encoding="utf-8"))
    k = yaml.safe_load((RULES_DIR / "checks.yaml").read_text(encoding="utf-8"))
    c, k = copy.deepcopy(c), copy.deepcopy(k)
    mutate(c, k)
    with pytest.raises(ValueError, match=msg):
        Rules(c, k)


def test_reports():
    r = analyse("1. Liability\nThe supplier is not liable for indirect loss.", RULES)
    md = to_markdown(r, RULES, "Test")
    assert "# Contract review: Test" in md and "Not legal advice" in md
    assert to_csv(r, RULES).splitlines()[0].endswith("clause_type,label,heading,text")


def test_consent_by_signing_is_still_flagged():
    assert "ai_training_on_data" in ids("The Customer consents to the Supplier using Customer Data to train its models.")
